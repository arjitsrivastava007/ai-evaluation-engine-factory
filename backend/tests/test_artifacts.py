"""Artifact parsers keep their kind-specific behavior and fail with a clear error."""

from io import BytesIO

import pytest
from docx import Document
from PIL import Image

from eval_factory.artifacts.ingest import ingest_bytes, kind_for_filename, safe_filename
from eval_factory.domain.models import ArtifactKind
from eval_factory.errors import FactoryError


def test_rejects_unknown_extension():
    with pytest.raises(FactoryError, match="Unsupported"):
        kind_for_filename("notes.exe")


def test_markdown_json_and_yaml_round_trip():
    markdown = ingest_bytes("design.md", b"# Retriever\n\nSearches the policy index.\n")
    assert markdown.kind is ArtifactKind.MARKDOWN
    assert "Retriever" in markdown.text

    parsed_json = ingest_bytes("arch.json", b'{"retriever": "faiss", "cite": true}')
    assert parsed_json.metadata["top_level_keys"] == ["retriever", "cite"]
    assert "faiss" in parsed_json.text

    parsed_yaml = ingest_bytes("arch.yaml", b"tools:\n  - search_policies\n")
    assert parsed_yaml.kind is ArtifactKind.YAML
    assert "search_policies" in parsed_yaml.text


def test_invalid_json_is_rejected():
    with pytest.raises(FactoryError, match="JSON"):
        ingest_bytes("broken.json", b"{not json")


def test_docx_includes_paragraphs_and_tables():
    document = Document()
    document.add_heading("Support bot", level=1)
    document.add_paragraph("Answers billing questions.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tool"
    table.rows[0].cells[1].text = "lookup_invoice"
    buffer = BytesIO()
    document.save(buffer)
    parsed = ingest_bytes("bot.docx", buffer.getvalue())
    assert parsed.kind is ArtifactKind.DOCX
    assert "billing questions" in parsed.text
    assert "lookup_invoice" in parsed.text


def test_image_records_dimensions_without_inventing_pixels():
    buffer = BytesIO()
    Image.new("RGB", (12, 8), color=(20, 40, 60)).save(buffer, format="PNG")
    parsed = ingest_bytes("diagram.png", buffer.getvalue())
    assert parsed.kind is ArtifactKind.IMAGE
    assert parsed.metadata["width"] == 12
    assert parsed.metadata["height"] == 8
    assert "12x8" in parsed.text


def test_pdf_page_count_is_recorded_even_when_text_is_absent():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    parsed = ingest_bytes("scan.pdf", buffer.getvalue())
    assert parsed.kind is ArtifactKind.PDF
    assert parsed.metadata["pages"] == 1


def test_safe_filename_strips_path_and_odd_characters():
    assert safe_filename("../../weird name!.md") == "weird_name_.md"
