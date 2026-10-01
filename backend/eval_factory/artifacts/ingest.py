"""Turn uploaded files into text the analyzer can read."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import yaml
from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from eval_factory.domain.models import ArtifactKind
from eval_factory.errors import FactoryError

_EXTENSIONS: dict[str, ArtifactKind] = {
    ".pdf": ArtifactKind.PDF,
    ".md": ArtifactKind.MARKDOWN,
    ".markdown": ArtifactKind.MARKDOWN,
    ".json": ArtifactKind.JSON,
    ".yaml": ArtifactKind.YAML,
    ".yml": ArtifactKind.YAML,
    ".docx": ArtifactKind.DOCX,
    ".png": ArtifactKind.IMAGE,
    ".jpg": ArtifactKind.IMAGE,
    ".jpeg": ArtifactKind.IMAGE,
    ".webp": ArtifactKind.IMAGE,
    ".gif": ArtifactKind.IMAGE,
    ".txt": ArtifactKind.TEXT,
}

_MEDIA_TYPES = {
    ArtifactKind.PDF: "application/pdf",
    ArtifactKind.MARKDOWN: "text/markdown",
    ArtifactKind.JSON: "application/json",
    ArtifactKind.YAML: "application/yaml",
    ArtifactKind.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ArtifactKind.IMAGE: "image/*",
    ArtifactKind.TEXT: "text/plain",
}

_TEXT_LIMIT = 200_000


class ParsedArtifact:
    """Extracted text plus the facts we can say for certain about the file."""

    def __init__(
        self,
        filename: str,
        kind: ArtifactKind,
        media_type: str,
        text: str,
        metadata: dict,
    ) -> None:
        self.filename = filename
        self.kind = kind
        self.media_type = media_type
        self.text = text
        self.metadata = metadata


def kind_for_filename(filename: str) -> ArtifactKind:
    suffix = Path(filename).suffix.lower()
    kind = _EXTENSIONS.get(suffix)
    if kind is None:
        supported = ", ".join(sorted({ext for ext in _EXTENSIONS}))
        raise FactoryError(
            f"Unsupported artifact type '{suffix or filename}'. Use one of: {supported}"
        )
    return kind


def safe_filename(filename: str) -> str:
    base = Path(filename).name.strip()
    cleaned = "".join(char if char.isalnum() or char in "._-" else "_" for char in base)
    cleaned = cleaned.strip("._") or "artifact"
    return cleaned[:180]


def ingest_bytes(filename: str, data: bytes) -> ParsedArtifact:
    if not data:
        raise FactoryError(f"{filename} is empty")
    kind = kind_for_filename(filename)
    try:
        text, metadata = _EXTRACTORS[kind](data)
    except FactoryError:
        raise
    except Exception as exc:
        raise FactoryError(f"Could not read {filename}: {exc}") from exc
    text = text.strip()
    if len(text) > _TEXT_LIMIT:
        text = text[:_TEXT_LIMIT]
        metadata["truncated"] = True
    if not text:
        text = f"No extractable text in {filename}."
    return ParsedArtifact(
        filename=safe_filename(filename),
        kind=kind,
        media_type=_MEDIA_TYPES[kind],
        text=text,
        metadata=metadata,
    )


def _read_text(data: bytes) -> tuple[str, dict]:
    text = data.decode("utf-8", errors="replace")
    return text, {"characters": len(text)}


def _read_json(data: bytes) -> tuple[str, dict]:
    raw = data.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise FactoryError(f"JSON artifact is not valid: {exc}") from exc
    pretty = json.dumps(parsed, indent=2, ensure_ascii=False, default=str)
    keys = list(parsed.keys()) if isinstance(parsed, dict) else []
    return pretty, {"top_level_keys": keys, "json_type": type(parsed).__name__}


def _read_yaml(data: bytes) -> tuple[str, dict]:
    raw = data.decode("utf-8", errors="replace")
    try:
        parsed = yaml.safe_load(raw)
    except yaml.YAMLError as exc:
        raise FactoryError(f"YAML artifact is not valid: {exc}") from exc
    pretty = yaml.safe_dump(parsed, sort_keys=False, allow_unicode=True) or ""
    keys = list(parsed.keys()) if isinstance(parsed, dict) else []
    return pretty, {"top_level_keys": keys}


def _read_pdf(data: bytes) -> tuple[str, dict]:
    reader = PdfReader(BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        pages.append(f"--- page {index} ---\n{extracted}".rstrip())
    return "\n\n".join(pages), {"pages": len(reader.pages)}


def _read_docx(data: bytes) -> tuple[str, dict]:
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
    }


def _read_image(data: bytes) -> tuple[str, dict]:
    try:
        image = Image.open(BytesIO(data))
        image.load()
    except UnidentifiedImageError as exc:
        raise FactoryError("Image artifact could not be decoded") from exc
    width, height = image.size
    metadata = {
        "format": image.format or "unknown",
        "width": width,
        "height": height,
        "mode": image.mode,
    }
    text = (
        f"Image artifact ({metadata['format']}, {width}x{height}, mode {image.mode}). "
        "Visual content is described by its file metadata. "
        "Pair a vision-capable provider when the picture itself must be judged."
    )
    return text, metadata


_EXTRACTORS = {
    ArtifactKind.PDF: _read_pdf,
    ArtifactKind.MARKDOWN: _read_text,
    ArtifactKind.JSON: _read_json,
    ArtifactKind.YAML: _read_yaml,
    ArtifactKind.DOCX: _read_docx,
    ArtifactKind.IMAGE: _read_image,
    ArtifactKind.TEXT: _read_text,
}
